"""Generate a realistic set of source workbooks for the treasury dashboard.

    python scripts/make_demo_data.py --out data/treasury

The point is to exercise the real path: these are *source* sheets (deals,
securities, clients), not pre-computed reports, and their column names are
deliberately the kind a bank actually exports — ``Cpty``, ``Notional Amt``,
``Mat. Date`` — so the mapping layer has something to do. One reported LCR
sheet is included to show the computed-vs-reported reconciliation.

Nothing here is real data.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import random
import sys

try:
    from openpyxl import Workbook
except ImportError:                                     # pragma: no cover
    print("needs openpyxl:  pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)

AS_OF = dt.date(2026, 7, 24)
BASE = "USD"

FX = {          # units of USD per 1 unit of currency
    "USD": 1.0, "EUR": 1.0850, "CHF": 1.1320, "GBP": 1.2740, "HKD": 0.1279,
    "CNH": 0.1382, "SGD": 0.7420, "JPY": 0.00655, "AUD": 0.6620, "AED": 0.2723,
}
# Annualised vols, written the way a rate column is read: a plain number in a
# rate column means percent (7.0 = 7%).
VOL = {"EUR": 7.0, "CHF": 7.5, "GBP": 8.0, "HKD": 1.2, "CNH": 4.5,
       "SGD": 5.0, "JPY": 10.0, "AUD": 11.0, "AED": 0.5}

CLIENTS = [
    # name, type, country, city, sector, rating, limit (USD), RM
    ("Pearl River Commercial Bank", "Bank", "China", "Guangzhou", "Banking", "BBB+", 120_000_000, "K. Lau"),
    ("Victoria Harbour Finance Ltd", "Financial institution", "Hong Kong", "Hong Kong", "Non-bank finance", "BBB", 80_000_000, "K. Lau"),
    ("Lion City Trading Pte", "Corporate", "Singapore", "Singapore", "Commodities", "BB+", 45_000_000, "M. Rossi"),
    ("Helvetia Precision AG", "Corporate", "Switzerland", "Zurich", "Industrials", "A-", 60_000_000, "M. Rossi"),
    ("Rhein Chemie GmbH", "Corporate", "Germany", "Frankfurt", "Chemicals", "BBB", 55_000_000, "M. Rossi"),
    ("Thames Asset Management", "Financial institution", "United Kingdom", "London", "Asset management", "A", 90_000_000, "J. Okafor"),
    ("Emirates Gateway Logistics", "Corporate", "UAE", "Dubai", "Logistics", "BB", 30_000_000, "J. Okafor"),
    ("Sakura Manufacturing KK", "Corporate", "Japan", "Tokyo", "Automotive", "A-", 40_000_000, "T. Berger"),
    ("Han River Steel Co", "Corporate", "Korea", "Seoul", "Metals", "BB+", 35_000_000, "T. Berger"),
    ("Formosa Semiconductor Supply", "Corporate", "Taiwan", "Taipei", "Technology", "BBB-", 25_000_000, "T. Berger"),
    ("Banco del Sur SA", "Bank", "Brazil", "Sao Paulo", "Banking", "BB", 20_000_000, "J. Okafor"),
    ("Alpine Private Bank", "Bank", "Switzerland", "Geneva", "Private banking", "A", 100_000_000, "M. Rossi"),
    ("Kowloon Textiles Holdings", "Corporate", "Hong Kong", "Kowloon", "Consumer", "BB-", 18_000_000, "K. Lau"),
    ("Sydney Infrastructure Trust", "Corporate", "Australia", "Sydney", "Infrastructure", "BBB+", 50_000_000, "T. Berger"),
    ("Grand Canal Shipping", "Corporate", "China", "Shanghai", "Shipping", "BB", 28_000_000, "K. Lau"),
    ("Marina Bay Wealth Partners", "Financial institution", "Singapore", "Singapore", "Wealth", "BBB+", 40_000_000, "M. Rossi"),
    ("Deutsche Mittelstand Bank AG", "Bank", "Germany", "Munich", "Banking", "A-", 75_000_000, "M. Rossi"),
    ("Gulf Petrochemical Trading", "Corporate", "UAE", "Abu Dhabi", "Energy", "BBB", 42_000_000, "J. Okafor"),
    ("Central Bank of the Republic", "Central bank", "Switzerland", "Bern", "Official", "AAA", 0, "Treasury"),
    ("Head Office — Group Treasury", "Internal", "Switzerland", "Zurich", "Intragroup", "", 0, "Treasury"),
]


def add(sheet, rows):
    for row in rows:
        sheet.append(row)


def styled(workbook, title, headers, note=""):
    sheet = workbook.create_sheet(title) if workbook.sheetnames != ["Sheet"] \
        else workbook.active
    sheet.title = title
    if note:
        sheet.append([note])
        sheet.append([])
    sheet.append(headers)
    return sheet


def d(days: int) -> dt.date:
    return AS_OF + dt.timedelta(days=days)


def usd_to(amount_usd, currency):
    """A USD-equivalent target expressed in the local currency."""
    return round(amount_usd / FX[currency], -3)


# Funding and asset blocks, sized in USD equivalents. Editing these numbers is
# how you make the demo bank look healthier or sicker.
FUNDING_BLOCKS = [
    # (label, product, cpty type, USD size, currencies, tenor days, flags)
    ("Corporate term deposits", "Time deposit", "Corporate", 520_000_000,
     ["USD", "USD", "EUR", "CHF", "HKD", "CNH", "SGD", "GBP"],
     [14, 30, 60, 90, 120, 180, 270, 365, 365], {}),
    ("Interbank taking", "MM borrowing", "Bank", 250_000_000,
     ["USD", "USD", "EUR", "CHF", "HKD"], [21, 30, 45, 60, 90, 150, 200], {}),
    ("Non-bank FI deposits", "Time deposit", "Financial institution", 90_000_000,
     ["USD", "EUR", "SGD"], [30, 60, 90, 150], {}),
    ("Operational balances", "Call account", "Financial institution", 90_000_000,
     ["USD", "HKD", "EUR"], [0], {"operational": True}),
]

ASSET_BLOCKS = [
    ("Corporate lending — short", "Loan", "Corporate", 190_000_000,
     ["USD", "USD", "EUR", "HKD", "CNH", "SGD"], [12, 25, 45, 90, 150, 300], {}),
    ("Corporate lending — term", "Loan", "Corporate", 590_000_000,
     ["USD", "USD", "EUR", "CHF", "HKD"], [430, 550, 730, 1100, 1460], {}),
    ("Interbank placements", "MM placement", "Bank", 145_000_000,
     ["USD", "HKD", "CHF", "EUR"], [3, 8, 15, 25, 45], {}),
]


def build_positions(rng, clients):
    """A balanced wholesale bank: funding blocks, asset blocks, then a hedge."""
    rows = []
    counter = {"n": 4100}

    def emit(product, name, kind, side, currency, amount, rate, start, tenor,
             flags=None):
        flags = flags or {}
        counter["n"] += 1
        prefix = {"Liability": "DEP", "Asset": "AST", "Capital": "CAP"}[side]
        maturity = "" if tenor is None else d(tenor).isoformat()
        rows.append([
            f"{prefix}{counter['n']}", flags.get("book", "Treasury"), name, kind,
            product, side, currency, round(amount, 2), round(rate, 3),
            d(start).isoformat(), maturity,
            "Y" if flags.get("operational") else "N",
            "Y" if flags.get("insured") else "N",
            "Y" if flags.get("secured") else "N",
            flags.get("collateral", ""),
            flags.get("country", ""),
        ])

    def pick(kinds):
        pool = [c for c in clients if c[1] in kinds]
        return rng.choice(pool)

    def spread_block(block, side):
        label, product, kinds, target, currencies, tenors, flags = block
        remaining = target
        while remaining > target * 0.03:
            share = min(remaining, target * rng.uniform(0.04, 0.13))
            remaining -= share
            currency = rng.choice(currencies)
            tenor = rng.choice(tenors)
            client = pick({kinds} if isinstance(kinds, str) else kinds)
            base_rate = {"USD": 4.3, "EUR": 2.6, "CHF": 1.0, "HKD": 3.9,
                         "CNH": 2.1, "SGD": 3.3, "GBP": 4.7}.get(currency, 3.0)
            margin = (rng.uniform(-0.35, 0.15) if side == "Liability"
                      else rng.uniform(0.9, 3.4))
            emit(product, client[0], client[1], side, currency,
                 usd_to(share, currency), max(base_rate + margin, 0.05),
                 -rng.randint(1, 200), tenor if tenor else None,
                 dict(flags, country=client[2]))

    for block in FUNDING_BLOCKS:
        spread_block(block, "Liability")
    for block in ASSET_BLOCKS:
        spread_block(block, "Asset")

    # Secured funding and intragroup term money.
    emit("Repo", "Pearl River Commercial Bank", "Bank", "Liability", "USD",
         55_000_000, 4.15, -3, 11, {"secured": True, "collateral": "L1"})
    emit("Reverse repo", "Deutsche Mittelstand Bank AG", "Bank", "Asset", "EUR",
         usd_to(30_000_000, "EUR"), 2.45, -2, 25,
         {"secured": True, "collateral": "L1"})
    emit("MM borrowing", "Head Office — Group Treasury", "Internal", "Liability",
         "USD", 120_000_000, 4.05, -120, 500)
    emit("Time deposit", "Alpine Private Bank", "Bank", "Liability", "CHF",
         usd_to(105_000_000, "CHF"), 1.35, -80, 620)

    # Shareholders' funds — 100% ASF, held in the reporting currency so it does
    # not masquerade as an FX position.
    emit("Capital", "Shareholders", "Internal", "Capital", "USD",
         250_000_000, 0.0, -3000, None)

    # Cash, central bank reserves and nostro working balances.
    for product, name, kind, currency, size in (
            ("Central bank", "Central Bank of the Republic", "Central bank", "CHF", 40_000_000),
            ("Central bank", "Central Bank of the Republic", "Central bank", "USD", 20_000_000),
            ("Nostro", "Pearl River Commercial Bank", "Bank", "CNH", 16_000_000),
            ("Nostro", "Victoria Harbour Finance Ltd", "Financial institution", "HKD", 18_000_000),
            ("Nostro", "Deutsche Mittelstand Bank AG", "Bank", "EUR", 12_000_000),
            ("Nostro", "Thames Asset Management", "Financial institution", "GBP", 9_000_000),
            ("Cash", "Head Office — Group Treasury", "Internal", "CHF", 5_000_000)):
        emit(product, name, kind, "Asset", currency, usd_to(size, currency), 0.0,
             -400, None, {"operational": product == "Nostro"})

    headers = ["Deal Ref", "Book", "Cpty", "Cpty Type", "Product", "A/L", "CCY",
               "Notional Amt", "Rate %", "Value Date", "Mat. Dt",
               "Operational?", "Insured?", "Secured?", "Collateral",
               "Country of Risk"]
    return headers, rows, emit


# A deliberate residual per currency so the FX panel has something real to show.
# CNH is left meaningfully open; everything else is close to square.
FX_RESIDUAL_USD = {"CNH": 17_500_000, "HKD": -4_200_000, "EUR": 3_100_000,
                   "CHF": -2_400_000, "GBP": 1_800_000, "SGD": -1_500_000}


def balance_currencies(rows, emit, securities, fx_trades, rng):
    """Close each currency's open position the way a treasury actually would.

    Whatever the deal book leaves open gets funded (or placed) in the same
    currency at 2-6 months, leaving a small deliberate residual so the FX panel
    has something real to show.
    """
    net = {}
    for row in rows:
        currency, amount, side = row[6], row[7], row[5]
        sign = 1.0 if side == "Asset" else -1.0
        net[currency] = net.get(currency, 0.0) + sign * amount * FX[currency]
    for security in securities:
        currency, mv = security[4], security[5]
        net[currency] = net.get(currency, 0.0) + mv * FX[currency]
    for trade in fx_trades:                    # buy leg long, sell leg short
        net[trade[4]] = net.get(trade[4], 0.0) + trade[5] * FX[trade[4]]
        net[trade[6]] = net.get(trade[6], 0.0) - trade[7] * FX[trade[6]]

    for currency in sorted(net):
        if currency == BASE:
            continue
        residual_target = FX_RESIDUAL_USD.get(currency,
                                              rng.uniform(-2_000_000, 2_000_000))
        gap = net[currency] - residual_target
        if abs(gap) < 1_000_000:
            continue
        tenor = rng.choice([65, 95, 130, 180])
        if gap > 0:            # long the currency: fund it locally
            emit("MM borrowing", "Alpine Private Bank", "Bank", "Liability",
                 currency, usd_to(gap, currency),
                 {"EUR": 2.5, "CHF": 1.0, "HKD": 3.9, "CNH": 2.1, "SGD": 3.3,
                  "GBP": 4.7, "JPY": 0.4, "AUD": 4.2, "AED": 4.4}.get(currency, 3.0),
                 -4, tenor)
        else:                  # short the currency: place the surplus
            emit("MM placement", "Deutsche Mittelstand Bank AG", "Bank", "Asset",
                 currency, usd_to(-gap, currency),
                 {"EUR": 2.4, "CHF": 0.9, "HKD": 3.8, "CNH": 2.0, "SGD": 3.2,
                  "GBP": 4.6, "JPY": 0.3, "AUD": 4.1, "AED": 4.3}.get(currency, 2.9),
                 -4, tenor)

    # Finally square the balance sheet itself in the reporting currency.
    total = 0.0
    for row in rows:
        total += (1.0 if row[5] == "Asset" else -1.0) * row[7] * FX[row[6]]
    for security in securities:
        total += security[5] * FX[security[4]]
    if total < 0:
        emit("MM placement", "Marina Bay Wealth Partners", "Financial institution",
             "Asset", BASE, -total, 4.2, -6, 75)
    elif total > 0:
        emit("MM borrowing", "Marina Bay Wealth Partners", "Financial institution",
             "Liability", BASE, total, 4.35, -6, 110)


def build_securities(rng):
    """Liquidity buffer plus a small investment tail. Sized in local currency."""
    plan = [
        # isin, name, issuer, type, ccy, USD size, rating, level, encumbered,
        # maturity days, coupon, risk weight
        ("CH0012345678", "Swiss Confederation 1.5% 2029", "Swiss Confederation",
         "Sovereign", "CHF", 62_000_000, "AAA", "L1", "N", 1180, 1.5, 0),
        ("US912828XX10", "US Treasury Note 4.0% 2027", "US Treasury", "Sovereign",
         "USD", 48_000_000, "AA+", "L1", "N", 520, 4.0, 0),
        ("US912828YY22", "US Treasury Bill 2026-08", "US Treasury", "Sovereign",
         "USD", 22_000_000, "AA+", "L1", "N", 21, 0.0, 0),
        ("DE0001102606", "Bund 2.3% 2031", "Bundesrepublik Deutschland", "Sovereign",
         "EUR", 24_000_000, "AAA", "L1", "Y", 1900, 2.3, 0),
        ("HK0000123456", "HKSAR Government Bond 2027", "HKSAR Government",
         "Sovereign", "HKD", 30_000_000, "AA+", "L1", "N", 430, 3.1, 0),
        ("XS2233445566", "EIB 2.8% 2028", "European Investment Bank", "PSE",
         "EUR", 26_000_000, "AAA", "L2A", "N", 900, 2.8, 20),
        ("XS9988776655", "Nestle 1.9% 2029", "Nestle SA", "Corporate", "CHF",
         14_000_000, "AA-", "L2A", "N", 1250, 1.9, 20),
        ("XS5566778899", "Financial Sub 5.4% 2028", "Thames Asset Management", "FI",
         "USD", 15_000_000, "BBB", "L2B", "N", 870, 5.4, 100),
        ("XS1122334455", "Gulf Petrochemical 6.2% 2030", "Gulf Petrochemical Trading",
         "Corporate", "USD", 24_000_000, "BB", "NONE", "N", 1500, 6.2, 100),
        ("XS7766554433", "Sydney Infrastructure 4.4% 2032", "Sydney Infrastructure Trust",
         "Corporate", "AUD", 16_000_000, "BBB+", "NONE", "N", 2200, 4.4, 100),
    ]
    rows = []
    for (isin, name, issuer, kind, currency, size_usd, rating, level,
         encumbered, days, coupon, risk_weight) in plan:
        market_value = usd_to(size_usd, currency)
        rows.append([isin, name, issuer, kind, currency, market_value,
                     round(market_value * rng.uniform(0.97, 1.01), -3), rating,
                     level, encumbered, d(days).isoformat(), coupon, risk_weight])
    headers = ["ISIN", "Security Name", "Issuer", "Issuer Type", "CCY",
               "Market Value", "Book Value", "Rating", "HQLA Level", "Encumbered",
               "Maturity", "Coupon %", "Risk Weight"]
    return headers, rows


def build_fx_trades(rng):
    rows = []
    trade = 800
    plan = [
        ("Spot", "USD", 25_000_000, "CHF", 22_100_000, 3),
        ("Forward", "EUR", 40_000_000, "USD", 43_400_000, 32),
        ("Forward", "USD", 30_000_000, "CNH", 217_000_000, 61),
        ("Swap", "HKD", 400_000_000, "USD", 51_150_000, 14),
        ("Forward", "USD", 18_000_000, "SGD", 24_250_000, 91),
        ("Forward", "JPY", 3_000_000_000, "USD", 19_650_000, 45),
        ("Forward", "USD", 12_000_000, "AUD", 18_130_000, 25),
        ("Spot", "GBP", 8_000_000, "USD", 10_190_000, 2),
        ("Forward", "AED", 55_000_000, "USD", 14_970_000, 75),
    ]
    for kind, buy_ccy, buy_amt, sell_ccy, sell_amt, days in plan:
        trade += 1
        rows.append([
            f"FX{trade}", kind, "FX Desk",
            rng.choice([c[0] for c in CLIENTS[:12]]),
            buy_ccy, buy_amt, sell_ccy, sell_amt,
            round(sell_amt / buy_amt, 6), d(-2).isoformat(), d(days).isoformat(),
        ])
    headers = ["Trade Ref", "Type", "Book", "Counterparty", "Buy CCY", "Buy Amount",
               "Sell CCY", "Sell Amount", "Deal Rate", "Trade Date", "Value Date"]
    return headers, rows


def build_commitments(rng):
    rows = []
    for name, kind, country, _city, _sector, _rating, limit, _rm in CLIENTS:
        if kind in {"Central bank", "Internal"} or not limit:
            continue
        if rng.random() < 0.45:
            continue
        facility = rng.choice(["Credit line", "Liquidity", "Trade finance", "Guarantee"])
        currency = rng.choice(["USD", "EUR", "HKD"])
        total = usd_to(rng.uniform(8_000_000, 32_000_000), currency)
        drawn = round(total * rng.uniform(0.0, 0.6), -3)
        rows.append([f"FAC-{abs(hash(name)) % 9000 + 1000}", name, kind, facility,
                     currency, total, drawn,
                     d(rng.randint(60, 900)).isoformat(), country])
    headers = ["Facility ID", "Client", "Client Type", "Line Type", "CCY",
               "Facility Limit", "Utilised", "Expiry", "Country"]
    return headers, rows


def credit_limits(position_rows, commitment_rows):
    """Set each client's limit from the exposure actually booked against them.

    Real limits are set first and drawn against second, but for a demo it has
    to be the other way round or half the book shows as a breach. One client is
    deliberately left over its line so the red state is visible.
    """
    exposure = {}
    for row in position_rows:
        if row[5] != "Asset":
            continue
        exposure[row[2]] = exposure.get(row[2], 0.0) + row[7] * FX[row[6]]
    for row in commitment_rows:
        undrawn = max(row[5] - row[6], 0.0)
        exposure[row[1]] = exposure.get(row[1], 0.0) + undrawn * FX[row[4]]

    limits = {}
    for name, used in exposure.items():
        headroom = 1.30 if name != BREACHED_CLIENT else 0.82
        limits[name] = max(round(used * headroom, -6), 5_000_000)
    return limits


BREACHED_CLIENT = "Grand Canal Shipping"


def build_clients(limits):
    rows = []
    for name, kind, country, city, sector, rating, _limit, rm in CLIENTS:
        rows.append([name, kind, sector, country, city, rating,
                     limits.get(name) or None, "USD", rm, ""])
    headers = ["Client Name", "Client Type", "Industry", "Country", "City",
               "Internal Rating", "Credit Limit", "Limit CCY",
               "Relationship Manager", "Remarks"]
    return headers, rows


def build_cashflows():
    rows = [
        [d(6).isoformat(), "CHF", -8_500_000, "Tax", "", "Quarterly corporate tax"],
        [d(12).isoformat(), "USD", -3_200_000, "Operating expense", "", "Payroll and premises"],
        [d(18).isoformat(), "EUR", 6_000_000, "Forecast client inflow",
         "Rhein Chemie GmbH", "Expected new term deposit"],
        [d(22).isoformat(), "USD", -25_000_000, "Capital", "", "Scheduled dividend to parent"],
        [d(30).isoformat(), "HKD", 40_000_000, "Fee income", "", "Trade finance commissions"],
        [d(9).isoformat(), "USD", 1_150_000, "Fee income", "", "FX spread and advisory fees"],
        [d(45).isoformat(), "CHF", -12_000_000, "Capex", "", "Core banking upgrade"],
    ]
    headers = ["Date", "CCY", "Amount", "Category", "Counterparty", "Description"]
    return headers, rows


def build_fx_rates():
    rows = []
    for currency, rate in FX.items():
        if currency == BASE:
            continue
        rows.append([currency + BASE, currency, BASE, rate, AS_OF.isoformat(),
                     VOL.get(currency, 0.08)])
    headers = ["Pair", "Currency", "Base", "Spot Rate", "Rate Date", "Annualised Vol"]
    return headers, rows


def build_market_rates(rng):
    rows = []
    curves = {"USD": 4.35, "EUR": 2.55, "CHF": 0.95, "HKD": 3.90, "CNH": 2.05,
              "GBP": 4.70, "SGD": 3.25}
    tenors = [("O/N", 0.0), ("1W", 0.02), ("1M", 0.05), ("2M", 0.08), ("3M", 0.12),
              ("6M", 0.20), ("1Y", 0.32)]
    for currency, anchor in curves.items():
        for tenor, bump in tenors:
            mid = anchor + bump
            rows.append([currency, tenor, round(mid - 0.05, 3), round(mid + 0.05, 3),
                         AS_OF.isoformat(), "AFS broker"])
    headers = ["CCY", "Tenor", "Bid", "Offer", "Quote Date", "Source"]
    return headers, rows


def build_reported_lcr():
    rows = [
        ["LCR", AS_OF.isoformat(), "Summary", "Liquidity Coverage Ratio", "", None,
         None, 1.34],
        ["LCR", AS_OF.isoformat(), "Summary", "Total HQLA", "USD", 640_000_000,
         None, 640_000_000],
        ["LCR", AS_OF.isoformat(), "Summary", "Net cash outflows", "USD",
         478_000_000, None, 478_000_000],
        ["NSFR", AS_OF.isoformat(), "Summary", "Net Stable Funding Ratio", "", None,
         None, 1.09],
    ]
    headers = ["Report", "As At", "Section", "Line Item", "CCY", "Amount",
               "Factor", "Weighted"]
    return headers, rows


def write_workbook(path, sheets):
    workbook = Workbook()
    first = True
    for title, note, headers, rows in sheets:
        sheet = workbook.active if first else workbook.create_sheet()
        sheet.title = title
        first = False
        if note:
            sheet.append([note])
            sheet.append([])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for index, header in enumerate(headers, start=1):
            width = max(11, min(34, len(str(header)) + 4))
            sheet.column_dimensions[sheet.cell(row=1, column=index)
                                    .column_letter].width = width
        sheet.freeze_panes = sheet.cell(row=(4 if note else 2), column=1)
    workbook.save(path)
    return path


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default="data/treasury", help="output folder")
    parser.add_argument("--seed", type=int, default=7)
    args = parser.parse_args(argv)
    rng = random.Random(args.seed)
    os.makedirs(args.out, exist_ok=True)

    position_headers, position_rows, emit = build_positions(rng, CLIENTS)
    security_headers, security_rows = build_securities(rng)
    fx_headers, fx_rows = build_fx_trades(rng)
    balance_currencies(position_rows, emit, security_rows, fx_rows, rng)
    commitment_headers, commitment_rows = build_commitments(rng)
    limits = credit_limits(position_rows, commitment_rows)

    written = []
    written.append(write_workbook(os.path.join(args.out, "positions.xlsx"), [
        ("Deals", f"Treasury deal blotter as at {AS_OF:%d %b %Y} — DEMO DATA",
         position_headers, position_rows),
    ]))
    written.append(write_workbook(os.path.join(args.out, "securities.xlsx"), [
        ("Portfolio", "Investment and liquidity portfolio — DEMO DATA",
         security_headers, security_rows),
    ]))
    written.append(write_workbook(os.path.join(args.out, "fx_trades.xlsx"), [
        ("FX Blotter", "", fx_headers, fx_rows),
    ]))
    written.append(write_workbook(os.path.join(args.out, "commitments.xlsx"), [
        ("Undrawn", "Committed but undrawn facilities",
         commitment_headers, commitment_rows),
    ]))
    written.append(write_workbook(os.path.join(args.out, "clients.xlsx"), [
        ("Client Master", "", *build_clients(limits)),
    ]))
    written.append(write_workbook(os.path.join(args.out, "cashflows.xlsx"), [
        ("Known Flows", "Anything the deal system doesn't know about",
         *build_cashflows()),
    ]))
    written.append(write_workbook(os.path.join(args.out, "market_data.xlsx"), [
        ("FX Rates", "", *build_fx_rates()),
        ("Money Market", "", *build_market_rates(rng)),
    ]))
    written.append(write_workbook(os.path.join(args.out, "reported_ratios.xlsx"), [
        ("Regulatory Return", "Last submitted return, for reconciliation",
         *build_reported_lcr()),
    ]))

    settings = {
        "_note": "Desk settings for the demo data set. Edit freely.",
        "entity": "Meridian Bank — Treasury",
        "base_currency": BASE,
        "as_of": AS_OF.isoformat(),
        "capital_base": 250_000_000,
        "nop_limit_pct_per_currency": 0.10,
        "nop_limit_pct_aggregate": 0.20,
        "concentration_top_n": 10,
        "cash_daily_days": 30,
        "lcr_table": "lcr_basel",
        "nsfr_table": "nsfr_basel",
    }
    settings_path = os.path.join(args.out, "settings.json")
    with open(settings_path, "w", encoding="utf-8") as fh:
        json.dump(settings, fh, indent=2)
    written.append(settings_path)

    for path in written:
        print(f"wrote {path}")
    print(f"\nnow run:  python dashboard.py --data {args.out} --open")
    return 0


if __name__ == "__main__":
    sys.exit(main())
