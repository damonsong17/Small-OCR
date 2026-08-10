"""Offline robustness tests -- stdlib unittest, no network, no Bloomberg.

    python -m unittest discover -s tests -v      (or: python tests/test_pipeline.py)

These lock in the behaviours that have actually broken before, so a regression
shows up on the offline machine instead of in a trading signal.
"""
from __future__ import annotations

import os
import sys
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from quote_ocr import sources                        # noqa: E402
from quote_ocr.arb import (                          # noqa: E402
    _is_ind,
    scan_across_channels,
    scan_surface_noarb,
)
from quote_ocr.bloomberg import (                    # noqa: E402
    FxPoint, all_pairs, order_pair, pip_of, quote_kind, to_outright,
)
from quote_ocr.docx_reader import parse_docx         # noqa: E402
from quote_ocr.implied import basis_of, implied_base_yield   # noqa: E402
from quote_ocr.parser import canonical_tenor          # noqa: E402
from quote_ocr.pricing import lookup_tenor            # noqa: E402


# --------------------------------------------------------------------------
class TestTenors(unittest.TestCase):
    def test_months_written_as_s(self):
        # the sheets write months as 1s/6s; downstream expects 1M/6M
        self.assertEqual(canonical_tenor("1s"), "1M")
        self.assertEqual(canonical_tenor("6s"), "6M")

    def test_overnight_variants(self):
        for raw in ("o/n", "o.n", "0/N"):
            self.assertEqual(canonical_tenor(raw), "O/N")

    def test_1y_and_12m_are_one_bucket(self):
        # a whole column silently vanished when these were treated as different
        self.assertEqual(lookup_tenor({"1Y": 0.042}, "12M"), 0.042)
        self.assertEqual(lookup_tenor({"12M": 0.042}, "1Y"), 0.042)

    def test_tenors_follow_the_data(self):
        surf = {"USD": {"bid": {"O/N": 0.04, "1W": 0.04, "2W": 0.04, "1M": 0.04,
                                "2M": 0.04, "3M": 0.04, "6M": 0.04, "1Y": 0.04},
                        "offer": {}}}
        self.assertEqual(sources.tenors_in(surf),
                         ["O/N", "1W", "2W", "1M", "2M", "3M", "6M", "1Y"])


# --------------------------------------------------------------------------
class TestCounterpartyAccess(unittest.TestCase):
    def test_no_kyc_names_are_excluded(self):
        for seg in ("Korean", "Taiwanese", "Indian", "ISLAMIC"):
            self.assertFalse(sources.is_tradeable(seg), seg)
            self.assertEqual(sources.access_status(seg), "no_kyc")

    def test_ocr_variants_of_those_names(self):
        for seg in ("ISLAMlC", "lndian", "Taiwanes", "ISLAMIC BANK", "korean "):
            self.assertFalse(sources.is_tradeable(seg), seg)

    def test_tradeable_and_blank_pass(self):
        self.assertTrue(sources.is_tradeable("Chinese"))
        self.assertTrue(sources.is_tradeable(""))

    def test_unknown_counterparty_excluded_by_default(self):
        self.assertEqual(sources.access_status("SomeNewBroker"), "unknown")
        self.assertFalse(sources.is_tradeable("SomeNewBroker"))


# --------------------------------------------------------------------------
class TestFxQuotes(unittest.TestCase):
    def test_points_vs_outright_classification(self):
        self.assertEqual(quote_kind("EUR/CNH 3M Curncy"), "outright")
        self.assertEqual(quote_kind("CNH+1M Curncy"), "outright")
        self.assertEqual(quote_kind("EURCNH3M Curncy"), "points")
        self.assertEqual(quote_kind("CGEU12M Curncy"), "points")

    def test_points_are_converted_to_an_outright(self):
        # EURCNH: spot 7.7972, 3M points -225.07  ->  7.7747
        got = to_outright(-225.0679, 7.7972, "points", 10000.0)
        self.assertAlmostEqual(got, 7.7972 - 225.0679 / 10000, places=9)

    def test_misclassified_points_are_caught(self):
        # declared outright but nowhere near spot -> treated as points
        got = to_outright(-225.0679, 7.7972, "outright", 10000.0)
        self.assertAlmostEqual(got, 7.7972 - 225.0679 / 10000, places=9)

    def test_pip_factor(self):
        self.assertEqual(pip_of("USDCNH"), 10000.0)
        self.assertEqual(pip_of("CNHJPY"), 100.0)

    def test_pair_direction_is_stable(self):
        self.assertEqual(order_pair("CNH", "USD"), "USDCNH")
        self.assertEqual(order_pair("USD", "EUR"), "EURUSD")
        self.assertEqual(order_pair("CNH", "HKD"), "HKDCNH")
        self.assertEqual(len(all_pairs(["USD", "CNH", "CHF", "EUR", "HKD"])), 10)


# --------------------------------------------------------------------------
class TestImpliedYield(unittest.TestCase):
    def test_round_trip(self):
        # a CIP-consistent forward must reverse-solve to the input rate
        S, r_q, r_b, act = 7.1850, 0.0150, 0.0430, 30
        F = S * (1 + r_q * act / 360) / (1 + r_b * act / 360)
        got = implied_base_yield("USDCNH", "1M", act, S, S, F, F, r_q)
        self.assertAlmostEqual(got.bid, r_b, places=9)

    def test_day_count_basis_per_currency(self):
        self.assertEqual(basis_of("CNH"), 360)   # confirmed on FXFA
        self.assertEqual(basis_of("HKD"), 365)
        self.assertEqual(basis_of("USD"), 360)

    def test_act_matters(self):
        # 1M is 33 actual days, not 30; using 30 shifts the yield materially
        S = 7.1850
        a = implied_base_yield("USDCNH", "1M", 33, S, S, 7.1725, 7.1725, 0.015)
        b = implied_base_yield("USDCNH", "1M", 30, S, S, 7.1725, 7.1725, 0.015)
        self.assertGreater(abs(a.bid - b.bid) * 1e4, 10)   # >10bp apart


# --------------------------------------------------------------------------
def _fx(spot, fwd, act=33):
    return FxPoint(tenor="3M", spot=spot, points=(fwd - spot) * 10000,
                   spot_bid=spot, spot_ask=spot, fwd_bid=fwd, fwd_ask=fwd, act=act)


class TestArbScan(unittest.TestCase):
    def setUp(self):
        self.fx = {"USDCNH": {"3M": _fx(7.1850, 7.1500), "1M": _fx(7.1850, 7.1725)}}

    def test_no_arb_surface_yields_nothing(self):
        # derive the CNH rate from CIP with the SAME act the scanner uses,
        # so the surface is exactly consistent -> no edge
        S, F, act, r_usd = 7.1850, 7.1500, 33, 0.0400
        r_cnh = ((F / S) * (1 + r_usd * act / 360) - 1) * 360 / act
        surf = {"USD": {"bid": {"3M": r_usd}, "offer": {"3M": r_usd}},
                "CNH": {"bid": {"3M": r_cnh}, "offer": {"3M": r_cnh}}}
        opps = scan_surface_noarb(surf, self.fx, ["USDCNH"], ["3M"],
                                  threshold_bps=5.0, allow_mismatch=False)
        self.assertEqual(opps, [])

    def test_obvious_arb_is_found(self):
        surf = {"USD": {"bid": {"3M": 0.0800}, "offer": {"3M": 0.0810}},
                "CNH": {"bid": {"3M": 0.0100}, "offer": {"3M": 0.0110}}}
        opps = scan_surface_noarb(surf, self.fx, ["USDCNH"], ["3M"],
                                  threshold_bps=1.0, allow_mismatch=False)
        self.assertTrue(opps)
        self.assertGreater(opps[0].pnl_bps, 0)

    def test_tenor_mismatch_is_labelled(self):
        surf = {"USD": {"bid": {"3M": 0.0800, "1M": 0.0800}, "offer": {"3M": 0.0810, "1M": 0.0810}},
                "CNH": {"bid": {"3M": 0.0100, "1M": 0.0100}, "offer": {"3M": 0.0110, "1M": 0.0110}}}
        opps = scan_surface_noarb(surf, self.fx, ["USDCNH"], ["1M", "3M"],
                                  threshold_bps=1.0, allow_mismatch=True)
        mism = [o for o in opps if o.mismatch]
        self.assertTrue(mism, "expected at least one mismatched-tenor route")
        o = mism[0]
        self.assertIn("TENOR MISMATCH", o.detail)
        self.assertIn("gap/rollover", o.risk_type)
        self.assertNotEqual(o.fx_tenor, o.lend_tenor)
        # every leg must be stated
        for f in (o.borrow_ccy, o.borrow_tenor, o.fx_tenor, o.lend_ccy, o.lend_tenor):
            self.assertTrue(f)

    def test_mismatch_can_be_disabled(self):
        surf = {"USD": {"bid": {"3M": 0.0800, "1M": 0.0800}, "offer": {"3M": 0.0810, "1M": 0.0810}},
                "CNH": {"bid": {"3M": 0.0100, "1M": 0.0100}, "offer": {"3M": 0.0110, "1M": 0.0110}}}
        opps = scan_surface_noarb(surf, self.fx, ["USDCNH"], ["1M", "3M"],
                                  threshold_bps=1.0, allow_mismatch=False)
        self.assertFalse([o for o in opps if o.mismatch])

    def test_indicative_leg_is_flagged(self):
        surf = {"USD": {"bid": {"3M": 0.0800}, "offer": {"3M": 0.0810}},
                "CNH": {"bid": {"3M": 0.0100}, "offer": {"3M": 0.0110}}}
        # the profitable route here is borrow CNH -> swap -> lend USD, so mark
        # the CNH borrowing side as the one derived from a reference rate
        opps = scan_surface_noarb(surf, self.fx, ["USDCNH"], ["3M"],
                                  threshold_bps=1.0, allow_mismatch=False,
                                  indicative={("CNH", "offer", "3M")})
        used = [o for o in opps if o.borrow_ccy == "CNH"]
        self.assertTrue(used)
        self.assertTrue(all(o.indicative for o in used))
        self.assertIn("INDICATIVE", used[0].detail)
        # a route that does not touch the reference rate stays firm
        clean = [o for o in opps if o.borrow_ccy == "USD"]
        self.assertFalse([o for o in clean if o.indicative])

    def test_indicative_flag_survives_1y_12m_spelling(self):
        # the flag must not be lost because the surface says 1Y and the key 12M
        self.assertTrue(_is_ind({("USD", "offer", "12M")}, "USD", "offer", "1Y"))
        self.assertFalse(_is_ind({("USD", "bid", "12M")}, "USD", "offer", "1Y"))


class TestCrossChannelArb(unittest.TestCase):
    """Borrow from one source, lend to another -- merging surfaces hides this."""

    def setUp(self):
        self.fx = {"USDCNH": {"3M": _fx(7.1850, 7.1500)}}

    def test_cheap_internal_expensive_afs(self):
        channels = {
            "MP":  {"USD": {"bid": {"3M": 0.0380}, "offer": {"3M": 0.0390}}},
            "AFS": {"USD": {"bid": {"3M": 0.0410}, "offer": {"3M": 0.0420}}},
        }
        opps = scan_across_channels(channels, self.fx, ["USDCNH"], ["3M"],
                                    threshold_bps=1.0, allow_mismatch=False)
        same = [o for o in opps if o.kind == "cross_channel_same_ccy"]
        self.assertTrue(same)
        o = same[0]
        self.assertEqual((o.borrow_channel, o.lend_channel), ("MP", "AFS"))
        self.assertAlmostEqual(o.pnl_bps, 20.0, places=2)   # 4.10 bid - 3.90 offer
        self.assertIn("@MP", o.legs())
        self.assertIn("@AFS", o.legs())

    def test_merging_would_have_hidden_it(self):
        channels = {
            "MP":  {"USD": {"bid": {"3M": 0.0380}, "offer": {"3M": 0.0390}}},
            "AFS": {"USD": {"bid": {"3M": 0.0410}, "offer": {"3M": 0.0420}}},
        }
        merged = sources.merge_surfaces(*channels.values())
        # best-of keeps offer 3.90 and bid 4.10 in ONE surface, which the
        # single-surface scan has no way to express as a trade between sources
        self.assertAlmostEqual(merged["USD"]["offer"]["3M"], 0.0390)
        self.assertAlmostEqual(merged["USD"]["bid"]["3M"], 0.0410)

    def test_offer_only_channel_can_be_borrowed_from_but_not_placed_with(self):
        """MM quotes one number and it is the OFFER: firm, and one-directional.

        Borrowing from MM and placing at AFS is a real trade. The reverse --
        placing funds with MM -- has no price at all and must not be invented.
        """
        channels = {
            "MM":  {"USD": {"bid": {}, "offer": {"3M": 0.0385}, "mid": {}}},
            "AFS": {"USD": {"bid": {"3M": 0.0410}, "offer": {"3M": 0.0420}}},
        }
        opps = scan_across_channels(channels, self.fx, ["USDCNH"], ["3M"],
                                    threshold_bps=1.0, allow_mismatch=False)
        same = [o for o in opps if o.kind == "cross_channel_same_ccy"]
        self.assertEqual(len(same), 1, "exactly one direction should be tradeable")
        o = same[0]
        self.assertEqual((o.borrow_channel, o.lend_channel), ("MM", "AFS"))
        self.assertAlmostEqual(o.pnl_bps, 25.0, places=2)   # 4.10 bid - 3.85 offer
        # firm, because an offer is an executable price -- not indicative
        self.assertFalse(o.indicative)
        self.assertNotIn("INDICATIVE", o.detail)

    def test_reference_only_channel_is_flagged_indicative(self):
        mm = {"USD": {"bid": {}, "offer": {}, "mid": {"3M": 0.0380}}}
        widened, ind = sources.apply_reference_sides(mm, half_spread_bps=0.0)
        channels = {
            "MM": widened,
            "AFS": {"USD": {"bid": {"3M": 0.0410}, "offer": {"3M": 0.0420}}},
        }
        opps = scan_across_channels(channels, self.fx, ["USDCNH"], ["3M"],
                                    threshold_bps=1.0, allow_mismatch=False,
                                    indicative={"MM": ind})
        same = [o for o in opps if o.kind == "cross_channel_same_ccy"]
        self.assertTrue(same)
        self.assertTrue(same[0].indicative)
        self.assertIn("INDICATIVE", same[0].detail)


class TestReferenceSides(unittest.TestCase):
    def test_mid_never_becomes_a_free_two_way_price(self):
        surf = {"USD": {"bid": {}, "offer": {}, "mid": {"3M": 0.0400}}}
        out, ind = sources.apply_reference_sides(surf, half_spread_bps=5.0)
        self.assertAlmostEqual(out["USD"]["bid"]["3M"], 0.0400 - 0.0005)
        self.assertAlmostEqual(out["USD"]["offer"]["3M"], 0.0400 + 0.0005)
        self.assertEqual(ind, {("USD", "bid", "3M"), ("USD", "offer", "3M")})

    def test_a_real_quoted_side_is_never_overwritten(self):
        surf = {"USD": {"bid": {"3M": 0.0390}, "offer": {}, "mid": {"3M": 0.0400}}}
        out, ind = sources.apply_reference_sides(surf, half_spread_bps=5.0)
        self.assertAlmostEqual(out["USD"]["bid"]["3M"], 0.0390)   # firm side kept
        self.assertNotIn(("USD", "bid", "3M"), ind)
        self.assertIn(("USD", "offer", "3M"), ind)

    def test_source_surface_is_not_mutated(self):
        surf = {"USD": {"bid": {}, "offer": {}, "mid": {"3M": 0.0400}}}
        sources.apply_reference_sides(surf, half_spread_bps=5.0)
        self.assertEqual(surf["USD"]["bid"], {})
        self.assertEqual(surf["USD"]["offer"], {})


# --------------------------------------------------------------------------
DOCX_XML = """<?xml version="1.0"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body><w:tbl>
{rows}
</w:tbl></w:body></w:document>"""


def _row(*cells):
    tcs = "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in cells)
    return f"<w:tr>{tcs}</w:tr>"


def _make_docx(path, rows):
    xml = DOCX_XML.format(rows="".join(_row(*r) for r in rows))
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml)


class TestStoreRoundTrip(unittest.TestCase):
    """A one-sided rate must survive the database, not vanish on the way in."""

    def setUp(self):
        self.db = Path("_test_tmp.db")
        if self.db.exists():
            os.remove(self.db)

    def tearDown(self):
        if self.db.exists():
            os.remove(self.db)

    def test_mid_only_quote_survives_the_store(self):
        from quote_ocr.models import Quote
        from quote_ocr.store import QuoteStore
        rows = [Quote(currency="USD", tenor="3M", mid="3.85", date="2026-08-07")]
        with QuoteStore(str(self.db)) as s:
            s.replace_quotes("MM", "2026-08-07", rows)
            surf = sources.surface_from_store(
                s, "2026-08-07", ["USD"], verbose=False, source="MM")
        self.assertAlmostEqual(surf["USD"]["mid"]["3M"], 0.0385)
        self.assertEqual(surf["USD"]["bid"], {})     # never invented on the way in
        self.assertEqual(surf["USD"]["offer"], {})

    def test_old_database_without_the_mid_column_is_migrated(self):
        import sqlite3
        from quote_ocr.store import QuoteStore
        # a database written by the previous release: no 'mid' column at all
        conn = sqlite3.connect(self.db)
        conn.execute("CREATE TABLE quotes (source TEXT, date TEXT, segment TEXT, "
                     "currency TEXT, benchmark TEXT, benchmark_rate TEXT, "
                     "tenor TEXT, bid TEXT, offer TEXT, confidence REAL, "
                     "source_file TEXT, page INTEGER, raw TEXT, ingested_at TEXT)")
        conn.execute("INSERT INTO quotes(source, date, currency, tenor, bid, offer) "
                     "VALUES ('AFS','2026-08-07','USD','3M','4.10','4.20')")
        conn.commit()
        conn.close()
        with QuoteStore(str(self.db)) as s:   # opening must upgrade it in place
            surf = sources.surface_from_store(
                s, "2026-08-07", ["USD"], verbose=False)
        self.assertAlmostEqual(surf["USD"]["bid"]["3M"], 0.0410)


class TestDocxReader(unittest.TestCase):
    def setUp(self):
        self.tmp = Path("_test_tmp.docx")

    def tearDown(self):
        if self.tmp.exists():
            os.remove(self.tmp)

    def test_currency_columns_reference_rates(self):
        _make_docx(self.tmp, [
            ("MONEY MARKET REFERENCE RATES(%)",),
            ("Tenor", "USD", "EUR"),
            ("1M", "3.85", "2.40"),
            ("3M", "3.95", "2.50"),
            ("Money Market Reference is for T+0 value.",),
        ])
        surf, meta = parse_docx(str(self.tmp), verbose=False)
        # The MM email quotes ONE number per currency and that number is the
        # OFFER -- the rate we can borrow at. It is firm, so it goes on the
        # offer side; the missing bid correctly means we cannot place there.
        self.assertAlmostEqual(surf["USD"]["offer"]["1M"], 0.0385)
        self.assertAlmostEqual(surf["EUR"]["offer"]["3M"], 0.0250)
        self.assertEqual(surf["USD"]["bid"], {})     # never invented
        self.assertEqual(surf["USD"]["mid"], {})
        # both currency columns survive even when they hold identical numbers
        self.assertIn("EUR", surf)
        self.assertEqual(meta["settle"], "T+0")
        self.assertIn("USD", meta["one_sided"])
        self.assertEqual(meta["single_side"], "offer")

    def test_single_side_is_overridable_without_a_code_change(self):
        _make_docx(self.tmp, [
            ("Tenor", "USD"),
            ("1M", "3.85"),
        ])
        surf, meta = parse_docx(str(self.tmp), verbose=False, single_side="mid")
        self.assertAlmostEqual(surf["USD"]["mid"]["1M"], 0.0385)
        self.assertEqual(surf["USD"]["offer"], {})
        self.assertEqual(meta["single_side"], "mid")

    def test_identical_values_in_two_columns(self):
        _make_docx(self.tmp, [
            ("Tenor", "USD", "EUR"),
            ("1M", "1.11", "1.11"),
        ])
        surf, _ = parse_docx(str(self.tmp), verbose=False)
        self.assertIn("USD", surf)
        self.assertIn("EUR", surf)   # de-duping on text used to drop this

    def test_bid_offer_layout(self):
        _make_docx(self.tmp, [
            ("USD RATES (%)",),
            ("Tenor", "BID", "OFFER"),
            ("1M", "3.80", "3.90"),
            ("note: T+2",),
        ])
        surf, meta = parse_docx(str(self.tmp), verbose=False)
        self.assertAlmostEqual(surf["USD"]["bid"]["1M"], 0.0380)
        self.assertAlmostEqual(surf["USD"]["offer"]["1M"], 0.0390)
        self.assertEqual(meta["settle"], "T+2")

    def test_bond_section_is_not_funding(self):
        _make_docx(self.tmp, [
            ("Tenor", "USD"),
            ("1M", "3.85"),
            ("USD SENIOR BOND",),
            ("Tenor", "REFERENCE RATES(%)"),
            ("3Y", "4.50"),
            ("5Y", "4.70"),
        ])
        surf, meta = parse_docx(str(self.tmp), verbose=False)
        self.assertIn("1M", surf["USD"]["offer"])
        self.assertNotIn("3Y", surf["USD"]["offer"])   # bond rows must not leak
        self.assertNotIn("5Y", surf["USD"]["offer"])
        self.assertTrue(meta["skipped"])


class TestTextSources(unittest.TestCase):
    """Broker lists, chat scrapes and PDF-to-Markdown reach the store too."""

    def setUp(self):
        self.tmp = Path("_test_quotes.txt")

    def tearDown(self):
        if self.tmp.exists():
            os.remove(self.tmp)

    def test_markdown_table_parses(self):
        # a PDF converted to Markdown: splitting on whitespace would make '|'
        # a column and shift every field by one
        self.tmp.write_text(
            "| currency | tenor | bid  | offer |\n"
            "|----------|-------|------|-------|\n"
            "| USD      | 3M    | 4.02 | 4.09  |\n", encoding="utf-8")
        surf = sources.surface_from_text(str(self.tmp))
        self.assertAlmostEqual(surf["USD"]["bid"]["3M"], 0.0402)
        self.assertAlmostEqual(surf["USD"]["offer"]["3M"], 0.0409)

    def test_csv_and_whitespace_forms_agree(self):
        self.tmp.write_text("USD,3M,4.02,4.09\n", encoding="utf-8")
        a = sources.surface_from_text(str(self.tmp))
        self.tmp.write_text("USD  3M  4.02  4.09\n", encoding="utf-8")
        b = sources.surface_from_text(str(self.tmp))
        self.assertEqual(a, b)

    def test_text_file_becomes_store_rows(self):
        from quote_ocr.batch import _quotes_from_text
        self.tmp.write_text("currency,tenor,bid,offer\nUSD,3M,4.02,4.09\n",
                            encoding="utf-8")
        rows = _quotes_from_text(str(self.tmp))
        self.assertEqual(len(rows), 1)
        q = rows[0]
        # same schema as an OCR or .docx row, so a query need not know the origin
        self.assertEqual((q.currency, q.tenor, q.bid, q.offer), ("USD", "3M", "4.02", "4.09"))
        self.assertEqual(q.mid, "")

    def test_unhandled_file_types_are_reported_not_skipped(self):
        import io
        from contextlib import redirect_stdout
        from quote_ocr.batch import ingest
        d = Path("_test_inbox/SRC")
        d.mkdir(parents=True, exist_ok=True)
        (d / "SRC_20260810.txt").write_text("USD,3M,4.02,4.09\n", encoding="utf-8")
        (d / "notes.rtf").write_text("not a quote file", encoding="utf-8")
        buf = io.StringIO()
        try:
            with redirect_stdout(buf):
                ingest(str(d.parent), "_test_out", make_xlsx=False)
            out = buf.getvalue()
            self.assertIn("ignored 1 file", out)
            self.assertIn("notes.rtf", out)
        finally:
            import shutil
            shutil.rmtree("_test_inbox", ignore_errors=True)
            shutil.rmtree("_test_out", ignore_errors=True)


try:
    import openpyxl
    _HAVE_XLSX = True
except ImportError:                                   # pragma: no cover
    _HAVE_XLSX = False


def _make_ftp_book(path, rows, fmt="0.00%"):
    """Minimal FTP sheet: Tenor | Offer Side (USD CNH) | Bid Side (USD CNH)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"], ws["C2"] = "Date:", "2026-08-10"
    ws["B3"], ws["C3"], ws["E3"] = "Tenor", "Offer Side", "Bid Side"
    for col, v in zip("CDEF", ["USD", "CNH", "USD", "CNH"]):
        ws[f"{col}4"] = v
    for i, (tenor, vals) in enumerate(rows):
        r = 5 + i
        ws[f"B{r}"] = tenor
        for col, v in zip("CDEF", vals):
            cell = ws[f"{col}{r}"]
            cell.value = v
            cell.number_format = fmt
    wb.save(path)


@unittest.skipUnless(_HAVE_XLSX, "openpyxl not installed")
class TestFtpSheet(unittest.TestCase):
    """The published FTP workbook: reading it, and filling it in."""

    def setUp(self):
        self.tmp = Path("_test_ftp.xlsx")
        self.out = Path("_test_ftp_out.xlsx")

    def tearDown(self):
        for f in (self.tmp, self.out):
            if f.exists():
                os.remove(f)

    def test_percent_formatted_cells_are_already_fractions(self):
        from quote_ocr.ftp_sheet import read_ftp_sheet
        # Excel stores 3.85% as 0.0385 when the cell is percent-formatted
        _make_ftp_book(self.tmp, [("3M", [0.0414, 0.0168, 0.0400, 0.0146])])
        surf, _ = read_ftp_sheet(str(self.tmp), verbose=False)
        self.assertAlmostEqual(surf["USD"]["offer"]["3M"], 0.0414)
        self.assertAlmostEqual(surf["CNH"]["bid"]["3M"], 0.0146)

    def test_bare_numbers_are_read_as_percent(self):
        from quote_ocr.ftp_sheet import read_ftp_sheet
        # the same rates typed into General-formatted cells as 4.14 etc.
        _make_ftp_book(self.tmp, [("3M", [4.14, 1.68, 4.00, 1.46])], fmt="General")
        surf, _ = read_ftp_sheet(str(self.tmp), verbose=False)
        self.assertAlmostEqual(surf["USD"]["offer"]["3M"], 0.0414)
        self.assertAlmostEqual(surf["CNH"]["bid"]["3M"], 0.0146)

    def test_placeholders_are_not_rates(self):
        from quote_ocr.ftp_sheet import read_ftp_sheet
        # an unfilled template must never read as 0% -- that would look like a
        # free deposit and manufacture arbitrage against every other currency
        _make_ftp_book(self.tmp, [("3M", ["x.xx%", "-", "x.xx%", "-"])])
        surf, meta = read_ftp_sheet(str(self.tmp), verbose=False)
        self.assertEqual(surf["USD"]["offer"], {})
        self.assertEqual(surf["CNH"]["bid"], {})
        self.assertEqual(len(meta["blank"]), 4)
        self.assertFalse(meta["unparsed"])

    def test_unreadable_cell_is_reported_not_dropped(self):
        from quote_ocr.ftp_sheet import read_ftp_sheet
        _make_ftp_book(self.tmp, [("3M", [0.04, "n/a", 0.039, 0.014])])
        surf, meta = read_ftp_sheet(str(self.tmp), verbose=False)
        self.assertNotIn("3M", surf["CNH"]["offer"])
        self.assertTrue(any("CNH 3M offer" in u for u in meta["unparsed"]))

    def test_write_then_read_round_trips(self):
        from quote_ocr.ftp_sheet import read_ftp_sheet, write_ftp_sheet
        _make_ftp_book(self.tmp, [("1M", ["x.xx%"] * 4), ("3M", ["x.xx%"] * 4)])
        surf = {"USD": {"offer": {"1M": 0.0392, "3M": 0.0414},
                        "bid": {"1M": 0.0378, "3M": 0.0400}},
                "CNH": {"offer": {"3M": 0.0168}, "bid": {"3M": 0.0146}}}
        write_ftp_sheet(surf, str(self.tmp), str(self.out), verbose=False)
        back, _ = read_ftp_sheet(str(self.out), verbose=False)
        self.assertAlmostEqual(back["USD"]["offer"]["3M"], 0.0414)
        self.assertAlmostEqual(back["CNH"]["bid"]["3M"], 0.0146)
        # a tenor we had no rate for keeps the template marker, never 0
        self.assertNotIn("1M", back["CNH"]["offer"])

    def test_1y_surface_fills_the_sheets_12m_row(self):
        from quote_ocr.ftp_sheet import read_ftp_sheet, write_ftp_sheet
        _make_ftp_book(self.tmp, [("12M", ["x.xx%"] * 4)])
        write_ftp_sheet({"USD": {"offer": {"1Y": 0.0438}, "bid": {}}},
                        str(self.tmp), str(self.out), verbose=False)
        back, _ = read_ftp_sheet(str(self.out), verbose=False)
        self.assertAlmostEqual(back["USD"]["offer"]["12M"], 0.0438)

    def test_writing_over_the_template_is_refused(self):
        from quote_ocr.ftp_sheet import write_ftp_sheet
        _make_ftp_book(self.tmp, [("3M", ["x.xx%"] * 4)])
        with self.assertRaises(ValueError):
            write_ftp_sheet({}, str(self.tmp), str(self.tmp), verbose=False)

    def test_an_empty_sheet_is_not_a_pass(self):
        """The worst possible bug in this tool: certifying a blank template."""
        import check_ftp
        _make_ftp_book(self.tmp, [("3M", ["x.xx%"] * 4)])
        rc = check_ftp.main([str(self.tmp)])
        self.assertEqual(rc, 2, "a sheet with no rates must not report 'no arbitrage'")

    def test_a_planted_arbitrage_is_caught_and_fails_the_run(self):
        import check_ftp
        # CNH bid set absurdly high: borrowing USD and placing CNH must profit
        _make_ftp_book(self.tmp, [("3M", [0.0414, 0.0168, 0.0400, 0.0310])])
        rc = check_ftp.main([str(self.tmp), "--tenors", "3M", "--ccy", "USD,CNH"])
        self.assertEqual(rc, 1, "a same-tenor round trip must fail the check")


if __name__ == "__main__":
    unittest.main(verbosity=2)
