"""Read .xlsx / .csv into plain tables, tolerating hand-made spreadsheets.

Treasury workbooks are written by people, not systems: a title row, a blank
line, merged cells, footnotes at the bottom, numbers stored as text with
thousands separators. This module's job is to find the real header row and hand
back ``list[dict]`` records; nothing downstream should know about cells.
"""
from __future__ import annotations

import csv
import datetime as _dt
import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set

# A cell that is one of these is treated as "no value" everywhere downstream.
BLANKS = {"", "-", "--", "n/a", "na", "nan", "none", "null", "#n/a"}

_NUM_RE = re.compile(r"^[\s]*\(?\s*[-+]?[\d,\s']*\.?\d+\s*\)?\s*%?$")
_DATE_FORMATS = (
    "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d.%m.%Y",
    "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y", "%Y%m%d", "%d-%b-%y",
)


@dataclass
class Table:
    """One rectangle of data pulled out of a sheet."""
    name: str                       # sheet name (or csv stem)
    source: str                     # file path it came from
    headers: List[str] = field(default_factory=list)
    records: List[Dict[str, Any]] = field(default_factory=list)
    header_row: int = 0             # 1-based, for error messages
    # Headers whose cells Excel formatted as percentages. Those values already
    # arrive as fractions (4.35% -> 0.0435), so a rate column must not be
    # divided by 100 a second time.
    percent_columns: Set[str] = field(default_factory=set)

    def __len__(self) -> int:
        return len(self.records)


# --- scalar coercion ----------------------------------------------------------
def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in BLANKS
    return False


def to_number(value: Any) -> Optional[float]:
    """Parse a spreadsheet cell as a number. ``(1,234)`` is negative; ``5%`` is 0.05."""
    if is_blank(value):
        return None
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not _NUM_RE.match(text):
        return None
    negative = text.startswith("(") and text.endswith(")")
    percent = text.endswith("%")
    cleaned = re.sub(r"[(),\s'%+]", "", text)
    try:
        number = float(cleaned)
    except ValueError:
        return None
    if percent:
        number /= 100.0
    return -number if negative else number


def to_date(value: Any) -> Optional[str]:
    """Return an ISO ``YYYY-MM-DD`` string, or None when the cell isn't a date."""
    if is_blank(value):
        return None
    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    text = str(value).strip()
    if re.fullmatch(r"\d{5}", text):        # Excel serial that arrived as text
        return excel_serial_to_date(int(text))
    for fmt in _DATE_FORMATS:
        try:
            return _dt.datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def excel_serial_to_date(serial: int) -> str:
    """Excel's 1900 date system (with its deliberate leap-year bug)."""
    return (_dt.date(1899, 12, 30) + _dt.timedelta(days=serial)).isoformat()


def to_bool(value: Any) -> Optional[bool]:
    if is_blank(value):
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"y", "yes", "true", "t", "1", "x"}:
        return True
    if text in {"n", "no", "false", "f", "0"}:
        return False
    return None


def to_text(value: Any) -> str:
    if is_blank(value):
        return ""
    if isinstance(value, (_dt.datetime, _dt.date)):
        return to_date(value) or ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# --- header detection ---------------------------------------------------------
def normalise_header(text: Any) -> str:
    """``"Notional (USD) "`` -> ``notional_usd`` so mapping can match loosely."""
    slug = re.sub(r"[^a-z0-9]+", "_", to_text(text).lower()).strip("_")
    return slug


def _looks_like_header(row: Sequence[Any]) -> int:
    """Score a row on how much it looks like a header (higher is better)."""
    filled = [c for c in row if not is_blank(c)]
    if len(filled) < 2:
        return -1
    wordy = sum(
        1 for c in filled
        if isinstance(c, str) and to_number(c) is None and to_date(c) is None
        and len(c.strip()) <= 40
    )
    if wordy < 2:
        return -1
    # Prefer wide, fully-worded rows; unique labels beat repeated ones.
    unique = len({normalise_header(c) for c in filled})
    return wordy * 2 + unique + len(filled)


def detect_header(rows: Sequence[Sequence[Any]], scan: int = 25) -> int:
    """Index of the most header-like row within the first ``scan`` rows."""
    best_index, best_score = -1, 0
    for index, row in enumerate(rows[:scan]):
        score = _looks_like_header(row)
        if score <= best_score:
            continue
        # A header must be followed by at least one row carrying data.
        following = rows[index + 1: index + 4]
        if not any(sum(0 if is_blank(c) else 1 for c in r) >= 2 for r in following):
            continue
        best_index, best_score = index, score
    return best_index


def _dedupe(headers: Sequence[str]) -> List[str]:
    seen: Dict[str, int] = {}
    out: List[str] = []
    for i, head in enumerate(headers):
        name = head or f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def table_from_rows(rows: Sequence[Sequence[Any]], name: str, source: str,
                    formats: Optional[Sequence[Sequence[str]]] = None) -> Optional[Table]:
    """Turn a raw grid into a Table, or None when there's no usable header."""
    header_index = detect_header(rows)
    if header_index < 0:
        return None
    headers = _dedupe([normalise_header(c) for c in rows[header_index]])
    percent = _percent_columns(headers, formats, header_index)
    records: List[Dict[str, Any]] = []
    for row in rows[header_index + 1:]:
        if all(is_blank(c) for c in row):
            continue
        record = {}
        for i, head in enumerate(headers):
            if not head.startswith("column_"):
                record[head] = row[i] if i < len(row) else None
        if any(not is_blank(v) for v in record.values()):
            records.append(record)
    return Table(name=name, source=source, headers=headers, records=records,
                 header_row=header_index + 1, percent_columns=percent)


def _percent_columns(headers: Sequence[str],
                     formats: Optional[Sequence[Sequence[str]]],
                     header_index: int) -> Set[str]:
    if not formats:
        return set()
    found: Set[str] = set()
    for column, header in enumerate(headers):
        seen = 0
        for row in formats[header_index + 1:]:
            if column >= len(row) or not row[column]:
                continue
            if "%" in row[column]:
                seen += 1
            elif seen:
                seen = 0
                break
            if seen >= 2:
                break
        if seen:
            found.add(header)
    return found


# --- file readers -------------------------------------------------------------
def read_csv(path: str) -> List[Table]:
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        rows = [row for row in csv.reader(fh)]
    name = os.path.splitext(os.path.basename(path))[0]
    table = table_from_rows(rows, name, path)
    return [table] if table else []


def read_xlsx(path: str) -> List[Table]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:      # pragma: no cover - dependency is in requirements
        raise RuntimeError(
            "reading .xlsx needs openpyxl: pip install openpyxl"
        ) from exc
    workbook = load_workbook(path, data_only=True, read_only=True)
    tables: List[Table] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            rows, formats = [], []
            for cells in sheet.iter_rows():
                rows.append([c.value for c in cells])
                formats.append([c.number_format for c in cells])
            table = table_from_rows(rows, sheet.title, path, formats)
            if table and table.records:
                tables.append(table)
    finally:
        workbook.close()
    return tables


def read_any(path: str) -> List[Table]:
    ext = os.path.splitext(path)[1].lower()
    if ext in {".xlsx", ".xlsm", ".xltx"}:
        return read_xlsx(path)
    if ext in {".csv", ".tsv", ".txt"}:
        return read_csv(path)
    return []


def walk_sources(root: str) -> List[str]:
    """Every readable data file under ``root`` (a file path is returned as-is)."""
    if os.path.isfile(root):
        return [root]
    found: List[str] = []
    for base, dirs, files in os.walk(root):
        dirs[:] = [d for d in dirs if not d.startswith((".", "~"))]
        for filename in sorted(files):
            if filename.startswith(("~$", ".")):
                continue
            if os.path.splitext(filename)[1].lower() in {
                ".xlsx", ".xlsm", ".xltx", ".csv", ".tsv"
            }:
                found.append(os.path.join(base, filename))
    return found


def read_tables(root: str) -> Iterable[Table]:
    for path in walk_sources(root):
        for table in read_any(path):
            yield table
