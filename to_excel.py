#!/usr/bin/env python3
"""Reshape a long-format quotes CSV back into the wide AFS-style Excel layout.

The pipeline emits one row per (tenor x currency); this rebuilds the original
two-table picture layout so you can eyeball it against the source image:

    python to_excel.py AFS_5th_exp.csv -o AFS_5th_exp.xlsx

Block A = rows carrying a benchmark (the "Chinese" USD/EUR/CNH/HKD grid),
columns grouped per currency as [benchmark, BID, OFFER].
Block B = the remaining market rows (Korean/Taiwanese/Indian/ISLAMIC), columns
grouped per market as [BID, OFFER]. Missing cells are shown as '-'.
"""
from __future__ import annotations

import argparse
import csv
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
GROUP_FILL = PatternFill("solid", fgColor="D9E1F2")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
BID_FILL = PatternFill("solid", fgColor="DDEBF7")
DASH = "-"


def _read(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def _ordered(values):
    return list(OrderedDict.fromkeys(values))


def build(rows):
    a_rows = [r for r in rows if (r.get("benchmark") or "").strip()]
    b_rows = [r for r in rows if not (r.get("benchmark") or "").strip()]

    wb = Workbook()
    ws = wb.active
    ws.title = "AFS"
    r = 1

    if a_rows:
        r = _block_a(ws, a_rows, r) + 2
    if b_rows:
        _block_b(ws, b_rows, r)

    ws.freeze_panes = "B1"
    ws.column_dimensions["A"].width = 8
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 11
    return wb


def _tenor_map(rows):
    """Ordered map of tenor -> display label, preserving file order.

    Uses the authoritative 'tenor' column (already normalised / OCR-corrected),
    not the verbatim 'raw' text which may contain the misread (e.g. '25').
    """
    m = OrderedDict()
    for row in rows:
        key = row["tenor"]
        if key and key not in m:
            m[key] = key
    return m


def _block_a(ws, rows, r0):
    currencies = _ordered(row["currency"] for row in rows)
    benchmark = {c: "" for c in currencies}
    data = {}  # (tenor, currency) -> (rate, bid, offer)
    for row in rows:
        c = row["currency"]
        benchmark[c] = benchmark[c] or row.get("benchmark", "")
        data[(row["tenor"], c)] = (
            row.get("benchmark_rate", ""), row.get("bid", ""), row.get("offer", "")
        )
    tenors = _tenor_map(rows)

    # Title banner.
    seg = next((row.get("segment") for row in rows if row.get("segment")), "")
    ncols = 1 + 3 * len(currencies)
    ws.cell(r0, 1, seg or "Table").font = Font(bold=True)
    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=ncols)
    ws.cell(r0, 1).alignment = CENTER

    # Currency group header row.
    gr = r0 + 1
    ws.cell(gr, 1, "").fill = GROUP_FILL
    col = 2
    for c in currencies:
        ws.merge_cells(start_row=gr, start_column=col, end_row=gr, end_column=col + 2)
        cell = ws.cell(gr, col, c)
        cell.font, cell.alignment, cell.fill = Font(bold=True), CENTER, GROUP_FILL
        col += 3

    # Sub-header row.
    sr = r0 + 2
    _hcell(ws, sr, 1, "Tenor")
    col = 2
    for c in currencies:
        _hcell(ws, sr, col, benchmark[c] or c)
        _hcell(ws, sr, col + 1, "BID", BID_FILL)
        _hcell(ws, sr, col + 2, "OFFER")
        col += 3

    # Data rows.
    rr = sr + 1
    for tnorm, disp in tenors.items():
        _dcell(ws, rr, 1, disp, bold=True)
        col = 2
        for c in currencies:
            rate, bid, offer = data.get((tnorm, c), ("", "", ""))
            _dcell(ws, rr, col, rate or DASH)
            _dcell(ws, rr, col + 1, bid or DASH, fill=BID_FILL)
            _dcell(ws, rr, col + 2, offer or DASH)
            col += 3
        rr += 1
    return rr - 1


def _block_b(ws, rows, r0):
    segments = _ordered(row["segment"] for row in rows)
    data = {}  # (tenor, segment) -> (currency, bid, offer)
    for row in rows:
        data[(row["tenor"], row["segment"])] = (
            row.get("currency", ""), row.get("bid", ""), row.get("offer", "")
        )
    tenors = _tenor_map(rows)

    gr = r0
    ws.cell(gr, 1, "").fill = GROUP_FILL
    col = 2
    for s in segments:
        ws.merge_cells(start_row=gr, start_column=col, end_row=gr, end_column=col + 1)
        cell = ws.cell(gr, col, s)
        cell.font, cell.alignment, cell.fill = Font(bold=True), CENTER, GROUP_FILL
        col += 2

    sr = r0 + 1
    _hcell(ws, sr, 1, "Tenor")
    col = 2
    for s in segments:
        bid_ccy = next(
            (data[(t, s)][0] for t in tenors if (t, s) in data and data[(t, s)][0]),
            "USD",
        )
        _hcell(ws, sr, col, f"{bid_ccy} BID", BID_FILL)
        _hcell(ws, sr, col + 1, "USD OFFER")
        col += 2

    rr = sr + 1
    for tnorm, disp in tenors.items():
        _dcell(ws, rr, 1, disp, bold=True)
        col = 2
        for s in segments:
            _, bid, offer = data.get((tnorm, s), ("", "", ""))
            _dcell(ws, rr, col, bid or DASH, fill=BID_FILL)
            _dcell(ws, rr, col + 1, offer or DASH)
            col += 2
        rr += 1
    return rr - 1


def _hcell(ws, r, c, val, fill=SUB_FILL):
    cell = ws.cell(r, c, val)
    cell.font, cell.alignment, cell.fill, cell.border = Font(bold=True), CENTER, fill, BORDER


def _dcell(ws, r, c, val, bold=False, fill=None):
    cell = ws.cell(r, c, val)
    cell.alignment, cell.border = CENTER, BORDER
    if bold:
        cell.font = Font(bold=True)
    if fill:
        cell.fill = fill


def _convert(csv_path: str, out: str) -> None:
    build(_read(csv_path)).save(out)
    print(f"Wrote {out}")


def main(argv=None):
    import glob
    import os

    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("csv", help="A quotes CSV, or a folder of CSVs to convert.")
    p.add_argument("-o", "--output", default=None,
                   help="Output .xlsx (single CSV) or output folder (for a dir).")
    args = p.parse_args(argv)

    if os.path.isdir(args.csv):
        out_dir = args.output or args.csv
        os.makedirs(out_dir, exist_ok=True)
        csvs = sorted(glob.glob(os.path.join(args.csv, "*.csv")))
        if not csvs:
            print(f"No .csv files in {args.csv}")
            return
        for c in csvs:
            stem = os.path.splitext(os.path.basename(c))[0]
            _convert(c, os.path.join(out_dir, stem + ".xlsx"))
    else:
        _convert(args.csv, args.output or (args.csv.rsplit(".", 1)[0] + ".xlsx"))


if __name__ == "__main__":
    main()
